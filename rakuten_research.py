#!/usr/bin/env python3
"""
楽天市場 競合分析ツール

自分の売れ筋商品名で検索し、競合商品の価格・レビュー・ポイント等を
分析してCSV/Excelに出力します。

使い方:
  python rakuten_research.py "商品名キーワード"
  python rakuten_research.py "商品名キーワード" --sort standard --pages 3
  python rakuten_research.py --file keywords.txt
"""

import argparse
import csv
import os
import sys
import time
from datetime import datetime
from pathlib import Path

import requests
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from tabulate import tabulate

load_dotenv()

RAKUTEN_API_URL = "https://app.rakuten.co.jp/services/api/IchibaItem/Search/20220601"


def get_app_id():
    app_id = os.getenv("RAKUTEN_APP_ID")
    if not app_id or app_id == "your_app_id_here":
        print("エラー: RAKUTEN_APP_ID が設定されていません。")
        print("1. https://webservice.rakuten.co.jp/ でアプリIDを取得してください（無料）")
        print("2. .env ファイルに RAKUTEN_APP_ID=取得したID を設定してください")
        sys.exit(1)
    return app_id


def search_items(keyword, app_id, sort="standard", page=1):
    """楽天市場APIで商品を検索する"""
    params = {
        "applicationId": app_id,
        "keyword": keyword,
        "sort": sort,
        "page": page,
        "hits": 30,
        "imageFlag": 1,
    }

    resp = requests.get(RAKUTEN_API_URL, params=params, timeout=15)
    resp.raise_for_status()
    return resp.json()


def extract_item_data(item_wrapper):
    """APIレスポンスから商品データを抽出する"""
    item = item_wrapper["Item"]

    # ポイント倍率の計算
    point_rate = item.get("pointRate", 1)
    price = item.get("itemPrice", 0)
    point_value = int(price * point_rate / 100)

    # 送料
    postage = "送料無料" if item.get("postageFlag", 0) == 0 else "送料別"
    # APIのpostageFlag: 0=送料込/無料, 1=送料別

    return {
        "商品名": item.get("itemName", ""),
        "ショップ名": item.get("shopName", ""),
        "価格": price,
        "送料": postage,
        "レビュー件数": item.get("reviewCount", 0),
        "レビュー平均": item.get("reviewAverage", "0"),
        "ポイント倍率": f"{point_rate}倍",
        "ポイント付与額": point_value,
        "実質価格": price - point_value,
        "商品URL": item.get("itemUrl", ""),
        "ショップコード": item.get("shopCode", ""),
        "商品コード": item.get("itemCode", ""),
    }


def analyze_competitors(items):
    """競合分析の統計を計算する"""
    if not items:
        return {}

    prices = [i["価格"] for i in items if i["価格"] > 0]
    real_prices = [i["実質価格"] for i in items if i["実質価格"] > 0]
    reviews = [i["レビュー件数"] for i in items]
    ratings = [float(i["レビュー平均"]) for i in items if float(i["レビュー平均"]) > 0]
    free_shipping = sum(1 for i in items if i["送料"] == "送料無料")

    return {
        "検索結果数": len(items),
        "価格_最安": min(prices) if prices else 0,
        "価格_最高": max(prices) if prices else 0,
        "価格_平均": int(sum(prices) / len(prices)) if prices else 0,
        "価格_中央値": sorted(prices)[len(prices) // 2] if prices else 0,
        "実質価格_最安": min(real_prices) if real_prices else 0,
        "レビュー件数_最大": max(reviews) if reviews else 0,
        "レビュー件数_平均": int(sum(reviews) / len(reviews)) if reviews else 0,
        "レビュー評価_平均": round(sum(ratings) / len(ratings), 2) if ratings else 0,
        "送料無料率": f"{free_shipping}/{len(items)} ({int(free_shipping/len(items)*100)}%)",
    }


def print_analysis(keyword, items, stats):
    """分析結果をコンソールに表示する"""
    print(f"\n{'='*60}")
    print(f"  検索キーワード: {keyword}")
    print(f"  検索日時: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    print(f"{'='*60}")

    # 統計サマリー
    print(f"\n【競合状況サマリー】")
    for key, val in stats.items():
        print(f"  {key}: {val:,}" if isinstance(val, int) else f"  {key}: {val}")

    # 上位商品一覧（簡易表示）
    table_data = []
    for i, item in enumerate(items[:10], 1):
        table_data.append([
            i,
            item["商品名"][:40] + ("..." if len(item["商品名"]) > 40 else ""),
            f"¥{item['価格']:,}",
            item["送料"],
            item["レビュー件数"],
            item["レビュー平均"],
            item["ポイント倍率"],
            f"¥{item['実質価格']:,}",
        ])

    headers = ["#", "商品名", "価格", "送料", "レビュー数", "評価", "PT倍率", "実質価格"]
    print(f"\n【上位10商品】")
    print(tabulate(table_data, headers=headers, tablefmt="grid"))

    # 売れてる理由の推測
    print(f"\n【売れている要因の分析】")
    if items:
        top = items[0]
        reasons = []
        if top["価格"] <= stats.get("価格_平均", 0):
            reasons.append("- 価格が平均以下で競争力がある")
        if top["レビュー件数"] >= stats.get("レビュー件数_平均", 0):
            reasons.append("- レビュー件数が多く信頼性が高い")
        if float(top["レビュー平均"]) >= 4.0:
            reasons.append("- レビュー評価が高い（4.0以上）")
        if top["送料"] == "送料無料":
            reasons.append("- 送料無料で購入ハードルが低い")
        if "倍" in top["ポイント倍率"] and top["ポイント倍率"] != "1倍":
            reasons.append("- ポイント倍率が高くお得感がある")
        if not reasons:
            reasons.append("- 商品タイトルやサムネイルの最適化が考えられます")
        for r in reasons:
            print(f"  {r}")


def save_to_excel(keyword, items, stats, output_dir="output"):
    """結果をExcelファイルに保存する"""
    Path(output_dir).mkdir(exist_ok=True)
    safe_keyword = keyword.replace("/", "_").replace("\\", "_")[:30]
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"{output_dir}/楽天競合分析_{safe_keyword}_{timestamp}.xlsx"

    wb = Workbook()

    # --- シート1: 商品一覧 ---
    ws1 = wb.active
    ws1.title = "商品一覧"

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    headers = list(items[0].keys()) if items else []
    for col, h in enumerate(headers, 1):
        cell = ws1.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font

    for row_idx, item in enumerate(items, 2):
        for col_idx, key in enumerate(headers, 1):
            ws1.cell(row=row_idx, column=col_idx, value=item[key])

    # 列幅の自動調整
    for col_idx, h in enumerate(headers, 1):
        ws1.column_dimensions[chr(64 + min(col_idx, 26))].width = max(len(h) * 2, 12)

    # --- シート2: 競合分析サマリー ---
    ws2 = wb.create_sheet("競合分析サマリー")
    ws2.cell(row=1, column=1, value="検索キーワード").font = Font(bold=True)
    ws2.cell(row=1, column=2, value=keyword)
    ws2.cell(row=2, column=1, value="検索日時").font = Font(bold=True)
    ws2.cell(row=2, column=2, value=datetime.now().strftime("%Y-%m-%d %H:%M:%S"))

    row = 4
    for key, val in stats.items():
        ws2.cell(row=row, column=1, value=key).font = Font(bold=True)
        ws2.cell(row=row, column=2, value=val)
        row += 1

    # --- シート3: 価格帯分析 ---
    ws3 = wb.create_sheet("価格帯分析")
    ws3.cell(row=1, column=1, value="価格帯").font = Font(bold=True)
    ws3.cell(row=1, column=2, value="商品数").font = Font(bold=True)
    ws3.cell(row=1, column=3, value="割合").font = Font(bold=True)

    prices = sorted([i["価格"] for i in items if i["価格"] > 0])
    if prices:
        # 価格帯を自動算出
        min_p, max_p = min(prices), max(prices)
        step = max((max_p - min_p) // 5, 1)
        ranges = []
        for start in range(min_p, max_p + 1, step):
            end = start + step - 1
            count = sum(1 for p in prices if start <= p <= end)
            if count > 0:
                ranges.append((f"¥{start:,}〜¥{end:,}", count))

        for row_idx, (range_label, count) in enumerate(ranges, 2):
            ws3.cell(row=row_idx, column=1, value=range_label)
            ws3.cell(row=row_idx, column=2, value=count)
            ws3.cell(row=row_idx, column=3, value=f"{count/len(prices)*100:.1f}%")

    wb.save(filename)
    return filename


def save_to_csv(keyword, items, output_dir="output"):
    """結果をCSVファイルに保存する"""
    Path(output_dir).mkdir(exist_ok=True)
    safe_keyword = keyword.replace("/", "_").replace("\\", "_")[:30]
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"{output_dir}/楽天競合分析_{safe_keyword}_{timestamp}.csv"

    if not items:
        return filename

    with open(filename, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=items[0].keys())
        writer.writeheader()
        writer.writerows(items)

    return filename


def run_research(keyword, app_id, sort="standard", pages=1):
    """1つのキーワードでリサーチを実行する"""
    all_items = []

    for page in range(1, pages + 1):
        try:
            data = search_items(keyword, app_id, sort=sort, page=page)
            items = [extract_item_data(i) for i in data.get("Items", [])]
            all_items.extend(items)

            total = data.get("count", 0)
            print(f"  ページ {page}: {len(items)}件取得 (全{total}件中)")

            if page < pages:
                time.sleep(1)  # APIレートリミット対策

        except requests.exceptions.HTTPError as e:
            print(f"  APIエラー (ページ{page}): {e}")
            break

    if not all_items:
        print(f"  「{keyword}」の検索結果が0件でした。")
        return

    stats = analyze_competitors(all_items)
    print_analysis(keyword, all_items, stats)

    excel_file = save_to_excel(keyword, all_items, stats)
    csv_file = save_to_csv(keyword, all_items)
    print(f"\n  Excel出力: {excel_file}")
    print(f"  CSV出力:   {csv_file}")

    return all_items


def main():
    parser = argparse.ArgumentParser(
        description="楽天市場 競合分析ツール - 自分の売れ筋商品の競合を分析します"
    )
    parser.add_argument("keyword", nargs="?", help="検索する商品名キーワード")
    parser.add_argument("--file", "-f", help="キーワードリストファイル（1行1キーワード）")
    parser.add_argument(
        "--sort", "-s",
        default="standard",
        choices=["standard", "-updateTimestamp", "-reviewCount", "+itemPrice", "-itemPrice"],
        help="ソート順: standard(標準), -reviewCount(レビュー多い順), +itemPrice(価格安い順), -itemPrice(価格高い順)",
    )
    parser.add_argument("--pages", "-p", type=int, default=2, help="取得ページ数 (1ページ=30件, デフォルト:2)")

    args = parser.parse_args()

    if not args.keyword and not args.file:
        parser.print_help()
        print("\n例: python rakuten_research.py \"ワイヤレスイヤホン\"")
        print("例: python rakuten_research.py --file keywords.txt --pages 3")
        sys.exit(1)

    app_id = get_app_id()

    keywords = []
    if args.file:
        with open(args.file, encoding="utf-8") as f:
            keywords = [line.strip() for line in f if line.strip()]
        print(f"キーワードファイルから {len(keywords)} 件のキーワードを読み込みました。")
    else:
        keywords = [args.keyword]

    print(f"\n楽天市場 競合分析ツール")
    print(f"ソート: {args.sort} / 取得ページ数: {args.pages}")

    for i, kw in enumerate(keywords):
        print(f"\n--- [{i+1}/{len(keywords)}] 「{kw}」を検索中 ---")
        run_research(kw, app_id, sort=args.sort, pages=args.pages)

        if i < len(keywords) - 1:
            time.sleep(1)

    print(f"\n完了！結果は output/ フォルダに保存されています。")


if __name__ == "__main__":
    main()
